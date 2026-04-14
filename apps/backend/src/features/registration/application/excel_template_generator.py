import io
import os
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

class ExcelTemplateGenerator:
    """
    Generates a dynamic Excel template for competitor registration,
    adding data validation (dropdowns) for Academies, Ranks, Sexes,
    and Special Conditions.
    """
    
    @staticmethod
    def generate_dynamic_template(academies: list[str], ranks: list[str], sexes: list[str]) -> io.BytesIO:
        # Load the base template
        current_dir = os.path.dirname(__file__)
        template_path = os.path.join(current_dir, "..", "api", "templates", "competitor_template.xlsx")
        
        # Load workbook
        wb = openpyxl.load_workbook(template_path)
        main_ws = wb.active
        
        # Create a hidden sheet to store the lists
        list_ws = wb.create_sheet(title="_Validaciones")
        list_ws.sheet_state = 'hidden'
        
        # Populate the lists into the hidden sheet
        # Column A: Sexes
        list_ws.cell(row=1, column=1, value="Sexos")
        for i, sex in enumerate(sexes, start=2):
            list_ws.cell(row=i, column=1, value=sex)
            
        # Column B: Ranks
        list_ws.cell(row=1, column=2, value="Rangos")
        for i, rank in enumerate(ranks, start=2):
            list_ws.cell(row=i, column=2, value=rank)
            
        # Column C: Academies
        list_ws.cell(row=1, column=3, value="Academias")
        for i, academy in enumerate(academies, start=2):
            list_ws.cell(row=i, column=3, value=academy)
            
        # Define the DataValidation ranges limit
        # In Excel, formulae can reference the ranges e.g. '_Validaciones!$A$2:$A$10'
        # To make it clean, we compute the exact rows used
        sex_range = f"'_Validaciones'!$A$2:$A${len(sexes) + 1}" if sexes else "A2:A2"
        rank_range = f"'_Validaciones'!$B$2:$B${len(ranks) + 1}" if ranks else "B2:B2"
        academy_range = f"'_Validaciones'!$C$2:$C${len(academies) + 1}" if academies else "C2:C2"
        
        # Create DataValidations
        # 1. Sex validation (Column G)
        dv_sex = DataValidation(type="list", formula1=sex_range, allow_blank=True)
        # 2. Rank validation (Column D)
        dv_rank = DataValidation(type="list", formula1=rank_range, allow_blank=True)
        # 3. Academy validation (Column I)
        dv_academy = DataValidation(type="list", formula1=academy_range, allow_blank=True)
        # 4. Special Condition validation (Column H) - Static SI/NO
        dv_condition = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
        
        # Add the validations to the main worksheet
        main_ws.add_data_validation(dv_sex)
        main_ws.add_data_validation(dv_rank)
        main_ws.add_data_validation(dv_academy)
        main_ws.add_data_validation(dv_condition)
        
        # Apply the validation to a large number of rows (e.g., from row 2 to 2000)
        # The template already has headers in row 1
        dv_sex.add("G2:G2000")
        dv_rank.add("D2:D2000")
        dv_academy.add("I2:I2000")
        dv_condition.add("H2:H2000")
        
        # Save to buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
